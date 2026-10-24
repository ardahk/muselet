"""Deterministic generator for the advisor-mode demo household (the Hendersons).

Run: python3 evals/fixtures/advisor/generate.py   (Python 3.12, stdlib + openpyxl)
Writes evals/fixtures/advisor/henderson/*. All truth values are recomputed from the written files by
the independent code in this file (not by importing brain.advisor); ml/tests/test_advisor.py then
checks that brain.advisor reproduces truth.json to the cent. The planted facts are asserted at the end.
Prices are fictional and fixed as of 2026-10-23 (the demo is offline).
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import zipfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parent / "henderson"
AS_OF = dt.date(2026, 10, 23)
CENT = Decimal("0.01")

TAXABLE = "Joint taxable brokerage"
TRAD = "Mark Traditional IRA"
ROTH = "Elena Roth IRA"
SAVINGS = "Savings"

US, INTL, BONDS, CASH = "US equity", "Intl equity", "Bonds", "Cash"

# Fictional fixed prices as of AS_OF.
SECURITIES = {
    "VTI": ("Vanguard Total Stock Market ETF", US, "328.50"),
    "VNQ": ("Vanguard Real Estate ETF", US, "86.40"),
    "AAPL": ("Apple Inc.", US, "252.30"),
    "MSFT": ("Microsoft Corp.", US, "489.60"),
    "PLTR": ("Palantir Technologies Inc.", US, "182.30"),
    "VXUS": ("Vanguard Total International Stock ETF", INTL, "71.20"),
    "BND": ("Vanguard Total Bond Market ETF", BONDS, "72.85"),
    "SWVXX": ("Schwab Value Advantage Money Fund", CASH, "1.00"),
    "VMFXX": ("Vanguard Federal Money Market Fund", CASH, "1.00"),
    "HYSA": ("High-yield savings account", CASH, "1.00"),
}

# (account, ticker, lot date or None, shares, cost per share)
LOTS = [
    (TAXABLE, "VTI", "2014-03-10", "330", "96.80"),
    (TAXABLE, "VTI", "2019-06-14", "726", "152.40"),
    (TAXABLE, "VTI", "2026-03-02", "40", "344.20"),      # small loss (< $1,000): not a TLH candidate
    (TAXABLE, "VNQ", "2022-01-03", "400", "108.90"),     # TLH ok
    (TAXABLE, "AAPL", "2012-09-17", "120", "24.60"),     # donation candidate
    (TAXABLE, "MSFT", "2015-11-02", "60", "53.80"),     # donation candidate
    (TAXABLE, "MSFT", "2024-07-15", "40", "455.00"),
    (TAXABLE, "PLTR", "2025-10-23", "100", "18.20"),     # held exactly 1 year: short-term, not a donation candidate
    (TAXABLE, "VXUS", "2017-05-08", "1200", "52.10"),
    (TAXABLE, "VXUS", "2026-02-09", "700", "76.40"),     # TLH, wash-sale blocked by the 2026-10-11 IRA buy
    (TAXABLE, "BND", "2013-04-15", "900", "70.10"),
    (TAXABLE, "BND", "2021-08-02", "1500", "87.40"),     # TLH ok
    (TAXABLE, "SWVXX", None, "57000.00", "1.00"),
    (TRAD, "VTI", "2016-01-19", "480", "104.70"),
    (TRAD, "VTI", "2026-09-30", "30", "331.10"),
    (TRAD, "VXUS", "2018-02-05", "2300", "58.40"),
    (TRAD, "VXUS", "2026-10-11", "100", "70.95"),
    (TRAD, "BND", "2016-01-19", "3200", "80.20"),
    (TRAD, "VMFXX", None, "18000.00", "1.00"),
    (ROTH, "VTI", "2017-03-06", "150", "121.30"),
    (ROTH, "VNQ", "2026-09-02", "60", "84.90"),
    (SAVINGS, "HYSA", None, "236480.00", "1.00"),
]

# Buys and sells across all accounts, 2026-08-24 .. 2026-10-23.
ACTIVITY = [
    ("2026-08-28", TAXABLE, "Sell", "MSFT", "25", "470.10"),
    ("2026-09-02", ROTH, "Buy", "VNQ", "60", "84.90"),
    ("2026-09-15", TAXABLE, "Sell", "MSFT", "25", "478.20"),
    ("2026-09-30", TRAD, "Buy", "VTI", "30", "331.10"),
    ("2026-10-11", TRAD, "Buy", "VXUS", "100", "70.95"),
]

TARGETS = {
    "as_of": AS_OF.isoformat(),
    "scope": "investment accounts (taxable + IRAs); the Savings account is the household liquidity bucket",
    "targets_pct": {US: "40", INTL: "20", BONDS: "35", CASH: "5"},
    "groups_pct": {"Equity": {"classes": [US, INTL], "target": "60"}},
    "rebalance_band_pp": "5",
}

PROFILE = {
    "household": "Henderson",
    "as_of": AS_OF.isoformat(),
    "members": [
        {"name": "Mark Henderson", "age": 62},
        {"name": "Elena Henderson", "age": 60},
    ],
    "filing_status": "married_filing_jointly",
    "state": "CA",
    "tax_year": 2026,
    "estimated_taxable_income": "150000.00",
    "estimated_taxable_income_note": "2026 federal taxable income before any Roth conversion",
    "monthly_expenses": "8000.00",
    "emergency_fund_months": 6,
    "goals": [
        "Retire fully when Mark turns 65",
        "Lake house down payment next spring (see last meeting notes)",
        "Start a charitable giving plan",
    ],
}

NOTES = """Henderson - meeting notes - 2026-07-23
Attendees: Mark Henderson, Elena Henderson, advisor

Life updates
- Mark plans to go part-time in January and retire fully at 65.
- Elena is consulting about 10 hours a week; income should be similar to last year.

Lake house
- They found a cabin near Lake Almanor. They want to put $120,000 down on the lake house next spring (target April 2027).
- Keep that money liquid and safe; they do not want it in the market.

Charitable giving
- Both are interested in charitable giving, starting this year: Elena mentioned the Alameda County food bank and the Sierra Club.
- Asked whether giving stock instead of cash makes sense. Follow up at the annual review.

Other
- Last year they gave $5,000 in cash to their church.
- Mark asked again about Roth conversions before RMDs start.
- Next meeting: annual review in October.
"""

ICS = "\r\n".join([
    "BEGIN:VCALENDAR",
    "VERSION:2.0",
    "PRODID:-//Muselet//Advisor demo//EN",
    "CALSCALE:GREGORIAN",
    "BEGIN:VEVENT",
    "UID:henderson-annual-review-20261026@muselet.local",
    "DTSTAMP:20261001T160000Z",
    "DTSTART:20261026T100000",
    "DTEND:20261026T110000",
    "SUMMARY:Henderson annual review",
    "LOCATION:Office",
    "DESCRIPTION:Annual review with Mark and Elena Henderson. Client folder: ~/Documents/Clients/Henderson",
    "END:VEVENT",
    "END:VCALENDAR",
]) + "\r\n"

HOLDINGS_HEADER = ["Account", "Ticker", "Name", "Asset class", "Lot date", "Shares", "Cost basis", "Price",
                   "Market value", "Unrealized G/L"]


def c(x: Decimal) -> Decimal:
    return x.quantize(CENT, rounding=ROUND_HALF_UP)


def s(x: Decimal) -> str:
    return str(c(x))


def write_files() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    rows = []
    for account, ticker, lot_date, shares, cps in LOTS:
        name, klass, price = SECURITIES[ticker]
        sh, price_d = Decimal(shares), Decimal(price)
        cost = c(sh * Decimal(cps))
        mv = c(sh * price_d)
        rows.append([account, ticker, name, klass, lot_date or "", shares, str(cost), price, str(mv), str(mv - cost)])

    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(HOLDINGS_HEADER)
    w.writerows(rows)
    (OUT / "portfolio.csv").write_text(buf.getvalue())

    wb = Workbook()
    ws = wb.active
    ws.title = "Holdings"
    ws.append(HOLDINGS_HEADER)
    for r in rows:
        ws.append([r[0], r[1], r[2], r[3],
                   dt.date.fromisoformat(r[4]) if r[4] else None,
                   float(r[5]), float(r[6]), float(r[7]), float(r[8]), float(r[9])])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        row[4].number_format = "yyyy-mm-dd"
        row[5].number_format = "#,##0.###"
        for cell in row[6:]:
            cell.number_format = "#,##0.00"
    for col, width in zip("ABCDEFGHIJ", (24, 8, 40, 12, 12, 11, 14, 10, 14, 14)):
        ws.column_dimensions[col].width = width
    fixed_time = dt.datetime(2026, 10, 23, 8, 0, 0)
    wb.properties.creator = "Muselet fixtures"
    wb.properties.created = fixed_time
    xbuf = io.BytesIO()
    wb.save(xbuf)
    # openpyxl stamps `modified` with now(); rewrite the zip with fixed metadata so output is byte-stable.
    src = zipfile.ZipFile(io.BytesIO(xbuf.getvalue()))
    stamp = fixed_time.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(OUT / "portfolio.xlsx", "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                start = data.index(b">", data.index(b"<dcterms:modified")) + 1
                end = data.index(b"</dcterms:modified>", start)
                data = data[:start] + stamp + data[end:]
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 10, 23, 8, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, data)

    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["Date", "Account", "Action", "Ticker", "Shares", "Price", "Amount"])
    for d, account, action, ticker, shares, price in ACTIVITY:
        w.writerow([d, account, action, ticker, shares, price, s(Decimal(shares) * Decimal(price))])
    (OUT / "activity_last_60d.csv").write_text(buf.getvalue())

    (OUT / "targets.json").write_text(json.dumps(TARGETS, indent=2) + "\n")
    (OUT / "profile.json").write_text(json.dumps(PROFILE, indent=2) + "\n")
    (OUT / "last_meeting_notes.txt").write_text(NOTES)
    (OUT / "annual_review.ics").write_bytes(ICS.encode())


# ---------------------------------------------------------------------------------------------------
# Truth: independent recomputation from the written files.

ACCOUNT_PREFERENCE = [TRAD, ROTH, TAXABLE]  # rebalance inside IRAs first (no realized gains)
PAIRS = {"VTI": "ITOT", "VXUS": "IXUS", "BND": "AGG", "VNQ": "SCHH", "VOO": "IVV", "VEA": "IEFA",
         "VWO": "IEMG", "BNDX": "IAGG"}
DEFAULT_TICKER = {US: "VTI", INTL: "VXUS", BONDS: "BND"}
BRACKET_22_TOP_MFJ_2026 = Decimal("211400")   # IRS Rev. Proc. 2025-32, sec. 4.01, Table 1
IRS_SOURCE = "https://www.irs.gov/pub/irs-drop/rp-25-32.pdf"


def pct(part: Decimal, total: Decimal) -> Decimal:
    return part * 100 / total


def lot_ref(i: int, r: dict) -> dict:
    return {"row": i, "account": r["Account"], "ticker": r["Ticker"], "lot_date": r["Lot date"],
            "shares": r["Shares"]}


def long_term(lot_date: dt.date) -> bool:
    try:
        anniversary = lot_date.replace(year=lot_date.year + 1)
    except ValueError:  # Feb 29
        anniversary = dt.date(lot_date.year + 1, 3, 1)
    return AS_OF > anniversary


def truth() -> dict:
    rows = list(csv.DictReader((OUT / "portfolio.csv").open()))
    activity = list(csv.DictReader((OUT / "activity_last_60d.csv").open()))
    targets = json.loads((OUT / "targets.json").read_text())
    profile = json.loads((OUT / "profile.json").read_text())
    notes = (OUT / "last_meeting_notes.txt").read_text()
    ics = (OUT / "annual_review.ics").read_text()

    for r in rows:
        assert Decimal(r["Market value"]) == c(Decimal(r["Shares"]) * Decimal(r["Price"]))
        assert Decimal(r["Unrealized G/L"]) == Decimal(r["Market value"]) - Decimal(r["Cost basis"])
    mv = [Decimal(r["Market value"]) for r in rows]

    by_account: dict[str, Decimal] = {}
    for r, v in zip(rows, mv):
        by_account[r["Account"]] = by_account.get(r["Account"], Decimal(0)) + v
    household_total = sum(mv)
    invested_idx = [i for i, r in enumerate(rows) if r["Account"] != SAVINGS]
    invested = sum(mv[i] for i in invested_idx)

    # Allocation
    tgt = {k: Decimal(v) for k, v in targets["targets_pct"].items()}
    band = Decimal(targets["rebalance_band_pp"])
    cls_val = {k: Decimal(0) for k in tgt}
    for i in invested_idx:
        cls_val[rows[i]["Asset class"]] += mv[i]
    classes = []
    for k in tgt:
        cur = pct(cls_val[k], invested)
        classes.append({"asset_class": k, "value": s(cls_val[k]), "current_pct": s(cur), "target_pct": s(tgt[k]),
                        "drift_pp": s(cur - tgt[k]), "out_of_band": abs(cur - tgt[k]) > band})
    groups = []
    for g, spec in targets["groups_pct"].items():
        gv = sum(cls_val[k] for k in spec["classes"])
        cur, t = pct(gv, invested), Decimal(spec["target"])
        groups.append({"group": g, "classes": spec["classes"], "value": s(gv), "current_pct": s(cur),
                       "target_pct": s(t), "drift_pp": s(cur - t), "out_of_band": abs(cur - t) > band})
    needed = any(x["out_of_band"] for x in classes)
    allocation = {"invested_total": s(invested), "band_pp": s(band), "classes": classes, "groups": groups,
                  "rebalance_needed": needed}

    # TLH (computed before rebalance: rebalance buys must avoid harvested tickers)
    min_loss, wash_days = Decimal(1000), 30
    tlh_lots = []
    for i, r in enumerate(rows, start=1):
        if r["Account"] != TAXABLE or not r["Lot date"]:
            continue
        gl = Decimal(r["Unrealized G/L"])
        if -gl < min_loss:
            continue
        lot_date = dt.date.fromisoformat(r["Lot date"])
        buys = [a for a in activity if a["Action"] == "Buy" and a["Ticker"] == r["Ticker"]
                and 0 <= (AS_OF - dt.date.fromisoformat(a["Date"])).days <= wash_days
                and not (a["Account"] == r["Account"] and a["Date"] == r["Lot date"] and a["Shares"] == r["Shares"])]
        entry = {"lot": lot_ref(i, r), "cost_basis": r["Cost basis"], "market_value": r["Market value"],
                 "loss": s(-gl), "term": "long" if long_term(lot_date) else "short",
                 "status": "wash_sale_blocked" if buys else "ok", "replacement": PAIRS.get(r["Ticker"])}
        if buys:
            last = max(dt.date.fromisoformat(a["Date"]) for a in buys)
            entry["blocking_buys"] = [{"date": a["Date"], "account": a["Account"], "ticker": a["Ticker"],
                                       "shares": a["Shares"]} for a in buys]
            entry["earliest_clean_sale"] = (last + dt.timedelta(days=wash_days + 1)).isoformat()
        tlh_lots.append(entry)
    tlh_lots.sort(key=lambda e: (-Decimal(e["loss"]), e["lot"]["row"]))
    harvest_total = sum(Decimal(e["loss"]) for e in tlh_lots if e["status"] == "ok")
    tlh = {"min_loss": s(min_loss), "wash_days": wash_days, "lots": tlh_lots,
           "harvestable_loss_total": s(harvest_total)}
    avoid = {e["lot"]["ticker"] for e in tlh_lots}

    # Rebalance
    trades = []
    post = dict(cls_val)
    if needed:
        delta = {k: c(invested * tgt[k] / 100 - cls_val[k]) for k in tgt if k != CASH}
        funds = {a: Decimal(0) for a in ACCOUNT_PREFERENCE}
        for i in invested_idx:
            if rows[i]["Asset class"] == CASH:
                funds[rows[i]["Account"]] += mv[i]
        for k in sorted((k for k in delta if delta[k] < 0), key=lambda k: delta[k]):
            need = -delta[k]
            for a in ACCOUNT_PREFERENCE:
                held: dict[str, Decimal] = {}
                for i in invested_idx:
                    if rows[i]["Account"] == a and rows[i]["Asset class"] == k:
                        held[rows[i]["Ticker"]] = held.get(rows[i]["Ticker"], Decimal(0)) + mv[i]
                for t in sorted(held, key=lambda t: (-held[t], t)):
                    if need <= 0:
                        break
                    amt = min(need, held[t])
                    assert a != TAXABLE, "fixture keeps rebalance sells inside the IRAs"
                    trades.append({"account": a, "side": "sell", "ticker": t, "asset_class": k, "amount": s(amt)})
                    need -= amt
                    funds[a] += amt
                    post[k] -= amt
        for k in sorted((k for k in delta if delta[k] > 0), key=lambda k: -delta[k]):
            need = delta[k]
            for a in ACCOUNT_PREFERENCE:
                if need <= 0 or funds[a] <= 0:
                    continue
                held = {}
                for i in invested_idx:
                    if rows[i]["Account"] == a and rows[i]["Asset class"] == k:
                        held[rows[i]["Ticker"]] = held.get(rows[i]["Ticker"], Decimal(0)) + mv[i]
                base = sorted(held, key=lambda t: (-held[t], t))[0] if held else DEFAULT_TICKER[k]
                ticker = PAIRS[base] if base in avoid else base
                amt = min(need, funds[a])
                tr = {"account": a, "side": "buy", "ticker": ticker, "asset_class": k, "amount": s(amt)}
                if ticker != base:
                    tr["instead_of"] = base
                trades.append(tr)
                need -= amt
                funds[a] -= amt
                post[k] += amt
        post[CASH] = invested - sum(post[k] for k in tgt if k != CASH)
    rebalance = {"rule": "full rebalance to target when any class is outside the band; IRAs first",
                 "trades": trades,
                 "post_trade_pct": {k: s(pct(post[k], invested)) for k in tgt},
                 "realized_gain": "0.00",
                 "total_bought": s(sum(Decimal(t["amount"]) for t in trades if t["side"] == "buy")),
                 "total_sold": s(sum(Decimal(t["amount"]) for t in trades if t["side"] == "sell"))}

    # Cash excess
    savings = by_account[SAVINGS]
    months = profile["emergency_fund_months"]
    monthly = Decimal(profile["monthly_expenses"])
    sentence = next(x for x in re.split(r"(?<=[.!?])\s+|\n", notes)
                    if "lake house" in x.lower() and "down" in x.lower())
    lake = Decimal(re.search(r"\$([\d,]+(?:\.\d\d)?)", sentence).group(1).replace(",", ""))
    excess = savings - months * monthly - lake
    cash = {"savings": s(savings), "months": months, "monthly_expenses": s(monthly),
            "emergency_reserve": s(months * monthly), "lake_house_reserve": s(lake),
            "excess": s(max(excess, Decimal(0)))}

    # Roth room
    ti = Decimal(profile["estimated_taxable_income"])
    roth = {"year": 2026, "filing_status": profile["filing_status"], "bracket_rate": "22%",
            "bracket_top": s(BRACKET_22_TOP_MFJ_2026), "taxable_income": s(ti),
            "room": s(max(BRACKET_22_TOP_MFJ_2026 - ti, Decimal(0))), "source": IRS_SOURCE}

    # Donation candidates
    cands = []
    for i, r in enumerate(rows, start=1):
        if r["Account"] != TAXABLE or not r["Lot date"]:
            continue
        gl, cost = Decimal(r["Unrealized G/L"]), Decimal(r["Cost basis"])
        if gl <= 0 or not long_term(dt.date.fromisoformat(r["Lot date"])):
            continue
        cands.append((gl / cost, i, r))
    cands.sort(key=lambda x: (-x[0], x[1]))
    donation_list = [{"lot": lot_ref(i, r), "cost_basis": r["Cost basis"], "market_value": r["Market value"],
                      "unrealized_gain": r["Unrealized G/L"], "gain_pct": s(g * 100)}
                     for g, i, r in cands[:2]]
    donation = {"n": 2, "candidates": donation_list,
                "market_value_total": s(sum(Decimal(d["market_value"]) for d in donation_list)),
                "gain_avoided_total": s(sum(Decimal(d["unrealized_gain"]) for d in donation_list))}

    # Meeting
    ev = dict(re.findall(r"^(SUMMARY|DTSTART|DTEND):(.*?)\r?$", ics, flags=re.M))
    iso = lambda v: dt.datetime.strptime(v, "%Y%m%dT%H%M%S").strftime("%Y-%m-%dT%H:%M")
    meeting = {"summary": ev["SUMMARY"], "start": iso(ev["DTSTART"]), "end": iso(ev["DTEND"])}

    opps = [
        {"key": "rebalance", "title": "Rebalance back to target", "impact_usd": rebalance["total_bought"],
         "impact_kind": "to move"},
        {"key": "cash_excess", "title": "Invest excess savings", "impact_usd": cash["excess"],
         "impact_kind": "idle in savings"},
        {"key": "roth_conversion", "title": "Roth conversion window", "impact_usd": roth["room"],
         "impact_kind": "convertible at 22% or less"},
        {"key": "tax_loss_harvest", "title": "Tax-loss harvesting", "impact_usd": tlh["harvestable_loss_total"],
         "impact_kind": "losses harvestable now"},
        {"key": "donate_appreciated", "title": "Gift appreciated stock to a donor-advised fund",
         "impact_usd": donation["gain_avoided_total"], "impact_kind": "capital gains avoided"},
    ]
    opps.sort(key=lambda o: (-Decimal(o["impact_usd"]), o["key"]))

    return {
        "household": profile["household"],
        "as_of": profile["as_of"],
        "meeting": meeting,
        "totals": {"household": s(household_total), "invested": s(invested),
                   "by_account": {a: s(v) for a, v in by_account.items()}},
        "allocation": allocation,
        "rebalance": rebalance,
        "tlh": tlh,
        "cash": cash,
        "roth": roth,
        "donation": donation,
        "opportunities": opps,
    }


def check_planted(t: dict) -> None:
    hh = Decimal(t["totals"]["household"])
    assert Decimal(1_400_000) <= hh <= Decimal(1_800_000), hh
    rows = list(csv.DictReader((OUT / "portfolio.csv").open()))
    positions = {(r["Account"], r["Ticker"]) for r in rows}
    assert 12 <= len(positions) <= 18, len(positions)
    multi = [p for p in positions if p[0] == TAXABLE and sum((r["Account"], r["Ticker"]) == p for r in rows) > 1]
    assert len(multi) >= 3, multi
    cls = {x["asset_class"]: x for x in t["allocation"]["classes"]}
    eq = t["allocation"]["groups"][0]
    assert abs(Decimal(eq["current_pct"]) - 68) < 1 and eq["out_of_band"], eq
    assert abs(Decimal(cls[BONDS]["current_pct"]) - 27) < 1 and cls[BONDS]["out_of_band"], cls[BONDS]
    assert abs(Decimal(cls[CASH]["current_pct"]) - 5) < 1, cls[CASH]
    assert all(t2["account"] != TAXABLE for t2 in t["rebalance"]["trades"])
    for k, v in t["rebalance"]["post_trade_pct"].items():
        assert abs(Decimal(v) - Decimal(TARGETS["targets_pct"][k])) < Decimal("0.01"), (k, v)
    lots = t["tlh"]["lots"]
    assert len(lots) == 3, lots
    blocked = [x for x in lots if x["status"] == "wash_sale_blocked"]
    assert len(blocked) == 1
    b = blocked[0]["blocking_buys"]
    assert len(b) == 1 and b[0]["account"] == TRAD and b[0]["ticker"] == blocked[0]["lot"]["ticker"]
    assert (AS_OF - dt.date.fromisoformat(b[0]["date"])).days == 12
    assert Decimal(t["cash"]["savings"]) > 2 * 6 * Decimal(PROFILE["monthly_expenses"])
    assert Decimal(t["cash"]["lake_house_reserve"]) == 120000 and Decimal(t["cash"]["excess"]) > 0
    assert t["roth"]["room"] == "61400.00"
    d = t["donation"]["candidates"]
    assert [x["lot"]["ticker"] for x in d] == ["AAPL", "MSFT"], d
    # the exactly-one-year PLTR lot would rank above MSFT if the > 1 year rule were ignored
    assert t["meeting"] == {"summary": "Henderson annual review", "start": "2026-10-26T10:00",
                            "end": "2026-10-26T11:00"}


def main() -> None:
    write_files()
    t = truth()
    check_planted(t)
    (OUT / "truth.json").write_text(json.dumps(t, indent=2, ensure_ascii=False) + "\n")
    a = {x["asset_class"]: x["current_pct"] for x in t["allocation"]["classes"]}
    print(f"household {t['totals']['household']}  invested {t['totals']['invested']}  allocation {a}")


if __name__ == "__main__":
    main()
